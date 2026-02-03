"""PowerPoint chunker with bounding box extraction."""

from pptx import Presentation
from loguru import logger

from src.models.entities.chunk_entities import BBox, ChunkCreate, SourceType
from src.modules.chunking.base_chunker import BaseChunker


class PPTXChunker(BaseChunker):
    """Chunk PowerPoint presentations with shape-level bbox."""

    def __init__(self, min_chunk_length: int = 20):
        """Initialize PPTX chunker."""
        self.min_chunk_length = min_chunk_length

    async def chunk(self, file_path: str, source_id: str, source_name: str) -> list[ChunkCreate]:
        """Extract chunks from PPTX with shape coordinates."""
        chunks = []

        try:
            prs = Presentation(file_path)

            for slide_num, slide in enumerate(prs.slides):
                for shape_idx, shape in enumerate(slide.shapes):
                    if not hasattr(shape, "text_frame"):
                        continue

                    text = shape.text.strip()
                    if not text or len(text) < self.min_chunk_length:
                        continue

                    # Convert EMUs (English Metric Units) to points
                    # 1 point = 12700 EMUs
                    emu_to_points = 12700

                    chunks.append(
                        ChunkCreate(
                            source_type=SourceType.UPLOADED_FILE,
                            source_id=source_id,
                            source_name=source_name,
                            content=text,
                            bbox=BBox(
                                page=slide_num + 1,
                                x=float(shape.left / emu_to_points),
                                y=float(shape.top / emu_to_points),
                                width=float(shape.width / emu_to_points),
                                height=float(shape.height / emu_to_points),
                            ),
                            metadata={
                                "slide_number": slide_num + 1,
                                "shape_index": shape_idx,
                                "shape_type": shape.shape_type.name if hasattr(shape.shape_type, "name") else "unknown",
                            },
                        )
                    )

            logger.info(f"Extracted {len(chunks)} chunks from PPTX: {source_name}")
            return chunks

        except Exception as e:
            logger.error(f"Error chunking PPTX {source_name}: {e}")
            raise
