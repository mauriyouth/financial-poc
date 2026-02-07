"""Test XLSX chunker with mock spreadsheet."""

import pytest
from openpyxl import Workbook

from src.models.entities.chunk_entities import SourceType
from src.modules.chunking.xlsx_chunker import XLSXChunker


@pytest.fixture
def mock_xlsx_path(tmp_path):
    """Create a mock Excel spreadsheet for testing."""
    xlsx_path = tmp_path / "test_spreadsheet.xlsx"

    wb = Workbook()

    # Sheet 1: Financial Data
    ws1 = wb.active
    ws1.title = "Financial Summary"

    ws1["A1"] = "Quarter"
    ws1["B1"] = "Revenue"
    ws1["C1"] = "Profit"
    ws1["D1"] = "Margin"

    ws1["A2"] = "Q1 2023"
    ws1["B2"] = 100
    ws1["C2"] = 25
    ws1["D2"] = "25%"

    ws1["A3"] = "Q2 2023"
    ws1["B3"] = 120
    ws1["C3"] = 32
    ws1["D3"] = "26.7%"

    ws1["A4"] = "Q3 2023"
    ws1["B4"] = 135
    ws1["C4"] = 38
    ws1["D4"] = "28.1%"

    # Sheet 2: Market Data
    ws2 = wb.create_sheet("Market Analysis")
    ws2["A1"] = "Region"
    ws2["B1"] = "Growth %"
    ws2["C1"] = "Market Share"

    ws2["A2"] = "North America"
    ws2["B2"] = "12%"
    ws2["C2"] = "35%"

    ws2["A3"] = "Europe"
    ws2["B3"] = "8%"
    ws2["C3"] = "28%"

    wb.save(str(xlsx_path))

    return str(xlsx_path)


@pytest.mark.asyncio
async def test_xlsx_chunker_basic(mock_xlsx_path):
    """Test basic XLSX chunking functionality."""
    chunker = XLSXChunker()

    chunks = await chunker.chunk(
        file_path=mock_xlsx_path, source_id="test_xlsx_001", source_name="test_spreadsheet.xlsx"
    )

    # Should extract chunks for each row with data
    assert len(chunks) > 0

    # Verify chunk properties
    assert chunks[0].source_type == SourceType.UPLOADED_FILE
    assert chunks[0].source_id == "test_xlsx_001"

    # Verify bbox with sheet and row info
    assert chunks[0].bbox is not None
    assert chunks[0].bbox.page is not None  # Sheet number
    assert chunks[0].bbox.y is not None  # Row number


@pytest.mark.asyncio
async def test_xlsx_chunker_sheet_metadata(mock_xlsx_path):
    """Test that sheet metadata is captured."""
    chunker = XLSXChunker()

    chunks = await chunker.chunk(file_path=mock_xlsx_path, source_id="test_xlsx_002", source_name="test.xlsx")

    # Verify metadata includes sheet information
    sheet_names = {chunk.metadata["sheet_name"] for chunk in chunks}
    assert "Financial Summary" in sheet_names
    assert "Market Analysis" in sheet_names

    # Verify row numbers
    for chunk in chunks:
        assert "row_number" in chunk.metadata
        assert chunk.metadata["row_number"] > 0


@pytest.mark.asyncio
async def test_xlsx_chunker_cell_joining(mock_xlsx_path):
    """Test that cells are properly joined with pipe delimiter."""
    chunker = XLSXChunker()

    chunks = await chunker.chunk(file_path=mock_xlsx_path, source_id="test_xlsx_003", source_name="test.xlsx")

    # Verify cells are joined with |
    assert any("|" in chunk.content for chunk in chunks)

    # Check specific content
    header_chunks = [c for c in chunks if "Quarter" in c.content]
    assert len(header_chunks) > 0
    assert "Revenue" in header_chunks[0].content


@pytest.mark.asyncio
async def test_xlsx_chunker_min_length(tmp_path):
    """Test minimum chunk length filtering."""
    xlsx_path = tmp_path / "short_sheet.xlsx"

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "A"
    ws["B1"] = "B"
    wb.save(str(xlsx_path))

    chunker = XLSXChunker(min_chunk_length=50)

    chunks = await chunker.chunk(file_path=str(xlsx_path), source_id="test_xlsx_004", source_name="short.xlsx")

    # Should filter out short rows
    assert len(chunks) == 0


@pytest.mark.asyncio
async def test_xlsx_chunker_empty_cells(tmp_path):
    """Test handling of empty cells."""
    xlsx_path = tmp_path / "sparse_sheet.xlsx"

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Data point one with sufficient length for testing purposes"
    ws["B1"] = None  # Empty cell
    ws["C1"] = "Data point two also with enough length to pass minimum"
    wb.save(str(xlsx_path))

    chunker = XLSXChunker()

    chunks = await chunker.chunk(file_path=str(xlsx_path), source_id="test_xlsx_005", source_name="sparse.xlsx")

    # Should handle empty cells gracefully
    assert len(chunks) > 0
    # Empty cells should not appear in content
    assert chunks[0].content.count("|") == 1  # Only one separator for 2 non-empty cells
