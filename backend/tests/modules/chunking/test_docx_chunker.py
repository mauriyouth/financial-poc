"""Test DOCX chunker with mock document."""

import pytest
from docx import Document

from src.models.entities.chunk_entities import SourceType
from src.modules.chunking.docx_chunker import DOCXChunker


@pytest.fixture
def mock_docx_path(tmp_path):
    """Create a mock Word document for testing."""
    docx_path = tmp_path / "test_document.docx"

    doc = Document()

    # Add title
    doc.add_heading("Annual Financial Report 2023", level=1)

    # Add paragraphs
    doc.add_paragraph(
        "The company delivered strong financial performance in 2023 with "
        "revenue growth of 12% and improved profitability across all segments."
    )

    doc.add_paragraph(
        "Key highlights include expansion into three new markets, "
        "successful product launches, and enhanced operational efficiency."
    )

    doc.add_heading("Market Overview", level=2)

    doc.add_paragraph(
        "Market conditions remained favorable throughout the year with "
        "sustained demand and positive industry trends supporting growth."
    )

    doc.save(str(docx_path))

    return str(docx_path)


@pytest.mark.asyncio
async def test_docx_chunker_basic(mock_docx_path):
    """Test basic DOCX chunking functionality."""
    chunker = DOCXChunker()

    chunks = await chunker.chunk(file_path=mock_docx_path, source_id="test_docx_001", source_name="test_document.docx")

    # Should extract chunks for paragraphs with sufficient length
    assert len(chunks) >= 3

    # Verify chunk properties
    assert chunks[0].source_type == SourceType.UPLOADED_FILE
    assert chunks[0].source_id == "test_docx_001"
    assert "financial performance" in chunks[0].content.lower()

    # Verify bbox with line numbers
    assert chunks[0].bbox is not None
    assert chunks[0].bbox.line_start is not None
    assert chunks[0].bbox.line_end is not None


@pytest.mark.asyncio
async def test_docx_chunker_paragraph_metadata(mock_docx_path):
    """Test that paragraph metadata is captured."""
    chunker = DOCXChunker()

    chunks = await chunker.chunk(file_path=mock_docx_path, source_id="test_docx_002", source_name="test.docx")

    # Verify metadata
    for chunk in chunks:
        assert "paragraph_number" in chunk.metadata
        assert "style" in chunk.metadata


@pytest.mark.asyncio
async def test_docx_chunker_min_length(tmp_path):
    """Test minimum chunk length filtering."""
    docx_path = tmp_path / "short_doc.docx"

    doc = Document()
    doc.add_paragraph("Short text here")
    doc.add_paragraph("Also short")
    doc.save(str(docx_path))

    chunker = DOCXChunker(min_chunk_length=50)

    chunks = await chunker.chunk(file_path=str(docx_path), source_id="test_docx_003", source_name="short.docx")

    # Should filter out short paragraphs
    assert len(chunks) == 0


@pytest.mark.asyncio
async def test_docx_chunker_sequential_line_numbers(mock_docx_path):
    """Test that line numbers are sequential."""
    chunker = DOCXChunker()

    chunks = await chunker.chunk(file_path=mock_docx_path, source_id="test_docx_004", source_name="test.docx")

    # Verify line numbers increase
    line_numbers = [chunk.bbox.line_start for chunk in chunks]
    assert line_numbers == sorted(line_numbers)
